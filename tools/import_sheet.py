"""Convertit le Google Sheet de la saison (export .xlsx) en fichier d'import pour l'Espace Suménium.

    python tools/import_sheet.py Saison_26.3.xlsx import.json [--connus noms.txt]

--connus : liste des pseudos réels (un par ligne, ex. la whitelist) pour corriger
           les fautes de frappe du sheet et signaler les pseudos introuvables.

Le JSON produit contient des données personnelles (tailles) : ne le commite pas,
importe-le depuis l'onglet Staff du site.
"""
import json
import re
import sys

import openpyxl

# Pseudos écrits autrement dans le sheet → pseudo Minecraft réel
ALIASES = {
    "roger rabit": "RogerRabit",
    "tassin": "Tassin98",
    "inomi": "InomiFraude",
    "lin_et_aru (linetaru)": "Lin_et_Aru",
    "linetaru": "Lin_et_Aru",
    "ayplex": "Ayplex_",
    "ldvdb": "LDVDB_",
    "pepemuntz": "PaPaMuntz13",
    "51percent/wayzie": "51Percent",
    "mekerrrr1012": "Mekerrr1012",
}
# Couleur de la case du joueur dans « Faction joueur » → activité prévue
ACTIVITE = {"FF3D85C6": "tres", "FF00FF00": "bcp", "FFFF9900": "moyen", "FFFF0000": "peu", "FFCC0000": "peu"}
MORT_COLS = ["oui", "sauf1", "sauf1_accord", "accord", "non"]
MC = re.compile(r"^[A-Za-z0-9_]{3,16}$")


def clean(v):
    if v is None:
        return ""
    return re.sub(r"[ \t]+", " ", str(v)).strip()


def main(src, dst, connus_path=None):
    wb = openpyxl.load_workbook(src, data_only=True)
    connus = {}
    if connus_path:
        for line in open(connus_path, encoding="utf-8"):
            if line.strip():
                connus[line.strip().lower()] = line.strip()

    def canon(name):
        n = clean(name)
        if not n:
            return None
        n = ALIASES.get(n.lower(), n)
        return connus.get(n.lower(), n)

    players = {}
    order = []

    def player(name):
        n = canon(name)
        if not n:
            return None
        key = n.lower()
        if key not in players:
            players[key] = {"mc_name": n}
            order.append(key)
        return players[key]

    # --- Factions et membres ---
    ws = wb["Faction joueur"]
    rows = {clean(ws.cell(r, 1).value).lower(): r for r in range(1, 40) if ws.cell(r, 1).value}
    factions = []
    for c in range(2, ws.max_column + 1):
        fname = clean(ws.cell(1, c).value)
        if not fname:
            continue
        sans = fname.lower() in ("sans faction", "slot faction")
        if not sans:
            def field(label):
                r = next((v for k, v in rows.items() if k.startswith(label)), None)
                return clean(ws.cell(r, c).value) if r else ""
            factions.append({
                "name": fname,
                "recrutement": field("recrutement"),
                "hub": field("projets de faction"),
                "projet_commu": field("projet communautaire"),
                "theme": field("thème de faction"),
                "autres_projets": field("autres projets"),
            })
        for r in range(2, 10):
            cell = ws.cell(r, c)
            if not cell.value:
                continue
            p = player(cell.value)
            if not sans:
                p["faction"] = fname
            act = ACTIVITE.get(cell.fill.fgColor.rgb if cell.fill.fgColor.type == "rgb" else "")
            if act:
                p["activite"] = act

    # --- Ordre du tirage des usines + choix ---
    ws = wb["Projets Usine"]
    headers = {c: clean(ws.cell(2, c).value) for c in range(3, ws.max_column + 1) if ws.cell(2, c).value}
    usine_choix = {}
    tirage = 0
    fnames = {f["name"].lower(): f for f in factions}
    for r in range(3, ws.max_row + 1):
        fname = clean(ws.cell(r, 2).value)
        if fname.lower() not in fnames:
            continue
        tirage += 1
        fnames[fname.lower()]["tirage"] = tirage
        picks = [re.sub(r"\s*\(.*\)$", "", headers[c]) for c in headers if headers[c] != "TOTAL" and ws.cell(r, c).value is True]
        if picks:
            usine_choix[fname] = picks

    # --- Métiers ---
    ws = wb["Métiers joueurs"]
    sans_joueur = []
    for r in range(2, ws.max_row + 1):
        name, metier, fac, mdesc, pdesc = (clean(ws.cell(r, c).value) for c in range(1, 6))
        if not (name or metier):
            continue
        if not name:
            sans_joueur.append(metier)
            continue
        p = player(name)
        if p.get("metier"):  # deux lignes pour le même joueur → on cumule
            p["metier"] += " / " + metier
            p["metier_desc"] = " / ".join(x for x in (p.get("metier_desc"), mdesc) if x)
        else:
            p["metier"], p["metier_desc"] = metier, mdesc
        if pdesc:
            p["perso_desc"] = pdesc
        if fac and not p.get("faction") and fac.lower() in fnames:
            p["faction"] = fnames[fac.lower()]["name"]

    # --- Mourir ou survivre ---
    ws = wb["Mourir ou Survivre "] if "Mourir ou Survivre " in wb.sheetnames else wb["Mourir ou Survivre"]
    for r in range(2, ws.max_row + 1):
        if clean(ws.cell(r, 1).value).lower().startswith("nombre total"):
            break
        for i, val in enumerate(MORT_COLS):
            v = clean(ws.cell(r, i + 1).value)
            if v and not v.lower().startswith("nombre total"):
                player(v)["mort"] = val

    # --- RP taille ---
    ws = wb["rp taille"]
    notes = []
    for r in range(2, ws.max_row + 1):
        name = clean(ws.cell(r, 1).value)
        m = ws.cell(r, 2).value
        if not name:
            continue
        p = player(name)
        if isinstance(m, (int, float)) and 0.5 <= m <= 2.5:
            p["taille_cm"] = round(m * 100)
        if clean(ws.cell(r, 5).value):
            notes.append(f"{p['mc_name']} veut changer de taille : {clean(ws.cell(r, 5).value)}")
        if clean(ws.cell(r, 7).value):
            p["taille_appliquee"] = round(m / 1.87, 4) if isinstance(m, (int, float)) and m else None

    out_players = [players[k] for k in order]
    data = {"factions": factions, "players": out_players, "usine_choix": usine_choix}
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)

    # --- Rapport ---
    print(f"{len(factions)} factions, {len(out_players)} joueurs, {sum(map(len, usine_choix.values()))} choix d'usines → {dst}")
    bad = [p["mc_name"] for p in out_players if not MC.match(p["mc_name"])]
    if bad:
        print("Pseudos invalides (ignorés à l'import) :", ", ".join(bad))
    if connus:
        absent = [p["mc_name"] for p in out_players if p["mc_name"].lower() not in connus]
        if absent:
            print("Pas encore vus sur le serveur (à vérifier) :", ", ".join(absent))
    if sans_joueur:
        print("Métiers sans joueur (non importés) :", ", ".join(sans_joueur))
    for n in notes:
        print("Note :", n)


if __name__ == "__main__":
    args = sys.argv[1:]
    connus = None
    if "--connus" in args:
        i = args.index("--connus")
        connus = args[i + 1]
        del args[i:i + 2]
    if len(args) != 2:
        sys.exit(__doc__)
    main(args[0], args[1], connus)
